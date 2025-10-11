# -*- coding: utf-8 -*-
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from matrix_view import MatrixView
from progress_dialog import ProgressDialog
from config import (ATTR_PERM, ATTR_SHORT, ATTR_OPTIONS,
                    same_crane_vert, same_crane_horz, is_crane_group)
from robot_helper import connect_robot, robot_delete_combos_by_names, robot_plan_numbers, robot_create_combo_with_number
from sap_helper import connect_sap2000, sap_delete_combos_by_names, sap_create_combo_and_add_factors
from config import ENGINE_ALREADY_CLEAN

# Excel support
_HAS_OPENPYXL = True
try:
    from openpyxl import Workbook, load_workbook
except Exception:
    _HAS_OPENPYXL = False

class ComboTab(ttk.Frame):
    def __init__(self, parent, get_groups_callable, get_context_callable):
        super().__init__(parent)
        self.get_groups = get_groups_callable
        self.get_context = get_context_callable
        self.combos=[]; self.columns_by_group=[]

        bar = ttk.Frame(self); bar.grid(row=0, column=0, sticky="ew", padx=8, pady=(8,4))
        bar.grid_columnconfigure(99, weight=1)

        ttk.Button(bar, text="Tổ hợp", command=self.build_combos).pack(side="left")
        ttk.Button(bar, text="Xuất Excel", command=self.export_excel).pack(side="left", padx=6)
        ttk.Button(bar, text="Nhập Excel", command=self.import_excel).pack(side="left", padx=6)

        ttk.Label(bar, text="Đưa vào:").pack(side="left", padx=(12,0))
        self.push_engine = tk.StringVar(value="ROBOT")
        ttk.Radiobutton(bar, text="ROBOT", value="ROBOT", variable=self.push_engine).pack(side="left", padx=(6,0))
        ttk.Radiobutton(bar, text="SAP2000", value="SAP2000", variable=self.push_engine).pack(side="left", padx=6)

        self.fast_mode = tk.BooleanVar(value=True)
        ttk.Checkbutton(bar, text="Chế độ nhanh", variable=self.fast_mode).pack(side="left", padx=(12,0))

        ttk.Button(bar, text="Đưa vào Engine", command=self.push_combos_to_engine).pack(side="right", padx=(0,6))
        ttk.Button(bar, text="Xóa combo trong Engine", command=self.delete_all_engine_combos).pack(side="right")

        self.matrix = MatrixView(self); self.matrix.grid(row=1, column=0, sticky="nsew", padx=8, pady=(0,8))
        self.grid_rowconfigure(1, weight=1); self.grid_columnconfigure(0, weight=1)

    # Máy sinh tổ hợp (1.0/0.9/0.7)
    def _generate(self, permanent, short_groups):
        import itertools
        temp_sets=[]
        for r in range(1, len(short_groups)+1):
            for chosen in itertools.combinations(short_groups, r):
                for pick in itertools.product(*chosen):
                    temp_sets.append(list(pick))
        if not temp_sets: temp_sets=[[]]
        combos=[]
        for temps in temp_sets:
            m=len(temps)
            if m==0 and not permanent: return []
            if m==0:
                assigns=[()]
            elif m==1:
                assigns=[((1.0, temps[0]),)]
            elif m==2:
                A,B=temps[0],temps[1]
                assigns=[((1.0,A),(0.9,B)), ((1.0,B),(0.9,A))]
            else:
                assigns=[]
                for i in range(m):
                    for j in range(m):
                        if i==j: continue
                        pack=[]
                        for k in range(m):
                            if k==i: pack.append((1.0, temps[k]))
                            elif k==j: pack.append((0.9, temps[k]))
                            else: pack.append((0.7, temps[k]))
                        assigns.append(tuple(pack))
            for a in assigns:
                parts=[]
                for (cid,nm) in permanent: parts.append((1.0, cid, nm))
                for (coef,(cid,nm)) in a: parts.append((coef, cid, nm))
                combos.append(parts)
        seen=set(); uniq=[]
        for parts in combos:
            key=tuple(sorted((p[1],p[0]) for p in parts))
            if key in seen: continue
            seen.add(key); uniq.append(parts)
        return uniq

    def build_combos(self):
        ctx = self.get_context() or {}
        groups = self.get_groups()
        if not groups:
            messagebox.showwarning("Chưa có nhóm","Hãy tạo nhóm & thêm tải ở tab LC, hoặc nhập Excel."); return

        ordered_names = list(groups.keys())

        permanent=[]; normal_short_groups=[]
        normal_columns_by_group=[]
        crane_v_cols=None; crane_h_cols=None

        crane_v_left_items=[];  crane_v_right_items=[]
        crane_h_left_items=[];  crane_h_right_items=[]

        for gname in ordered_names:
            raw_items=list(groups[gname].get("items") or [])  # (cid, nm) hoặc (cid, nm, side)
            cols=[(str(it[0]), str(it[1])) for it in raw_items]
            if not cols:
                normal_columns_by_group.append((gname, []))
                continue

            if same_crane_vert(gname):
                crane_v_cols = (gname, cols)
                for it in raw_items:
                    cid, nm = str(it[0]), str(it[1])
                    side = (it[2] if len(it)>=3 else "L")
                    (crane_v_left_items if side=="L" else crane_v_right_items).append((cid,nm))
            elif same_crane_horz(gname):
                crane_h_cols = (gname, cols)
                for it in raw_items:
                    cid, nm = str(it[0]), str(it[1])
                    side = (it[2] if len(it)>=3 else "L")
                    (crane_h_left_items if side=="L" else crane_h_right_items).append((cid,nm))
            else:
                normal_columns_by_group.append((gname, cols))
                attr=groups[gname].get("attr")
                if attr==ATTR_PERM:
                    permanent.extend([(str(it[0]), str(it[1])) for it in raw_items])
                elif attr==ATTR_SHORT:
                    normal_short_groups.append([(str(it[0]), str(it[1])) for it in raw_items])

        columns_by_group = list(normal_columns_by_group)
        if crane_v_cols is not None: columns_by_group.append(crane_v_cols)
        if crane_h_cols is not None: columns_by_group.append(crane_h_cols)

        generated_all = []

        def _gen_with(v_items, h_items):
            short_groups_all = list(normal_short_groups)
            if v_items: short_groups_all.append(v_items)
            if h_items: short_groups_all.append(h_items)
            g = self._generate(permanent, short_groups_all)
            if h_items:
                set_v = {cid for (cid,_nm) in v_items}
                set_h = {cid for (cid,_nm) in h_items}
                filtered=[]
                for parts in g:
                    has_v = any(str(p[1]) in set_v for p in parts)
                    has_h = any(str(p[1]) in set_h for p in parts)
                    if has_h and not has_v: continue  # nếu có NGANG thì bắt buộc có ĐỨNG cùng bên
                    filtered.append(parts)
                return filtered
            return g

        # Trái chỉ ghép Trái, Phải chỉ ghép Phải
        generated_all.extend(_gen_with(crane_v_left_items,  crane_h_left_items))
        generated_all.extend(_gen_with(crane_v_right_items, crane_h_right_items))

        seen=set(); generated=[]
        for parts in generated_all:
            key=tuple(sorted((p[1],p[0]) for p in parts))
            if key in seen: continue
            seen.add(key); generated.append(parts)

        keys=[k for _g, cols in columns_by_group for (k,_lab) in cols]
        rows=[]; self.combos=[]
        def fmt(v): return (str(v).rstrip('0').rstrip('.') if isinstance(v, float) else str(v))
        for i,parts in enumerate(generated, start=1):
            name=f"COMBO{i:03d}"; cmap={}
            for (coef,cid,_nm) in parts: cmap[str(cid)] = coef
            row=[name]+[ ("" if k not in cmap else fmt(cmap[k])) for k in keys ]
            rows.append(row); self.combos.append({"name":name,"map":cmap})
        self.columns_by_group = columns_by_group
        self.matrix.build(columns_by_group, rows)
        messagebox.showinfo("Xong", f"Đã sinh {len(rows)} tổ hợp (tách Trái/Phải, không có BOTH).")

    # Export/Import Excel
    def export_excel(self):
        if not _HAS_OPENPYXL:
            messagebox.showerror("Thiếu thư viện","Cần gói 'openpyxl' (pip install openpyxl)"); return
        if not self.combos:
            messagebox.showwarning("Chưa có tổ hợp","Bấm 'Tổ hợp' hoặc 'Nhập Excel' trước."); return
        col_ids=[]; col_names=[]
        for (_g, lst) in self.columns_by_group:
            for k, lab in lst:
                col_ids.append(str(k)); col_names.append(str(lab))
        path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                            filetypes=[("Excel Workbook","*.xlsx")],
                                            initialfile="combos.xlsx")
        if not path: return
        wb = Workbook(); ws = wb.active; ws.title = "Combos"
        ws.append(["COMBO"] + col_ids)
        ws.append([""] + col_names)
        for c in self.combos:
            ws.append([c["name"]] + [c["map"].get(k,"") for k in col_ids])
        try: wb.save(path)
        except Exception as e:
            messagebox.showerror("Lỗi ghi Excel", f"{e}"); return
        messagebox.showinfo("Xuất Excel", f"Đã lưu: {path}")

    def import_excel(self):
        if not _HAS_OPENPYXL:
            messagebox.showerror("Thiếu thư viện","Cần gói 'openpyxl' (pip install openpyxl)"); return
        path = filedialog.askopenfilename(filetypes=[("Excel Workbook","*.xlsx")])
        if not path: return
        try:
            wb = load_workbook(path, data_only=True); ws = wb.active
        except Exception as e:
            messagebox.showerror("Lỗi đọc Excel", f"{e}"); return

        excel_ids=[str(cell.value).strip() for cell in ws[1][1:] if cell.value not in (None,""," ")]
        excel_labels=[str(cell.value).strip() if cell.value not in (None,""," ") else "" for cell in ws[2][1:len(excel_ids)+1]]

        if not self.columns_by_group:
            cols=[(cid, (excel_labels[i] if i < len(excel_labels) else cid)) for i, cid in enumerate(excel_ids)]
            self.columns_by_group=[("TỪ EXCEL", cols)]

        imported=[]
        for r in ws.iter_rows(min_row=3, values_only=True):
            if r is None: continue
            name = str(r[0]).strip() if r[0] else ""
            if not name: continue
            cmap={}
            for idx, v in enumerate(r[1:]):
                if idx >= len(excel_ids): break
                key = excel_ids[idx]
                if v in (None,""," "): continue
                try: cmap[str(key)] = float(v)
                except Exception: continue
            imported.append({"name":name, "map":cmap})
        if not imported:
            messagebox.showwarning("Không có dữ liệu", "Không tìm thấy tổ hợp (từ hàng 3)."); return

        keys_current=[k for _g, lst in self.columns_by_group for (k,_lab) in lst]
        rows=[]; self.combos=[]
        def fmt(v): return (str(v).rstrip('0').rstrip('.')) if isinstance(v,float) else str(v)
        for c in imported:
            rows.append([c["name"]] + [ ("" if str(k) not in c["map"] else fmt(c["map"][str(k)])) for k in keys_current ])
            self.combos.append({"name": c["name"], "map": {str(k): float(c["map"][str(k)]) for k in c["map"]}})
        self.matrix.build(self.columns_by_group, rows)
        messagebox.showinfo("Nhập Excel", f"Đã nạp {len(self.combos)} tổ hợp từ Excel.")

    # Push/Delete
    def _get_case_name_resolver(self, ctx):
        raw = ctx.get("raw_by_id", {}) or {}
        def resolver(cid_str):
            d = raw.get(str(cid_str))
            if not d: return None
            return d.get("name")
        return resolver

    def push_combos_to_engine(self):
        if not self.combos:
            messagebox.showwarning("Chưa có tổ hợp","Hãy bấm 'Tổ hợp' hoặc 'Nhập Excel' trước."); return

        target = self.push_engine.get().upper()
        ctx = self.get_context() or {}
        prog = ProgressDialog(self, title=f"Đưa combo vào {target}", total=len(self.combos))
        prog.set_status("Chuẩn bị..."); prog.log(">> Begin")

        if target == "ROBOT":
            rob, _ = connect_robot(visible=not self.fast_mode.get())
            if not rob:
                prog.log("!! Không tìm thấy phiên Robot."); prog.close()
                messagebox.showerror("Không kết nối", "Không tìm thấy phiên Robot đang mở.")
                return
            cases = rob.Project.Structure.Cases

            if self.fast_mode.get():
                names_set = {c["name"] for c in self.combos if c.get("name")}
                if names_set:
                    prog.log(f"- Đang xóa trước combo trùng tên ({len(names_set)}) ...")
                removed = robot_delete_combos_by_names(cases, names_set, prog=prog if names_set else None)
                prog.log(f"- Đã xóa trước {removed} combo trùng tên (ROBOT).")

            plan = robot_plan_numbers(cases, len(self.combos))
            try: cases.BeginMultiOperation()
            except Exception: pass

            try:
                created = 0
                for idx, (number, combo) in enumerate(zip(plan, self.combos), start=1):
                    if prog.cancelled: break
                    name = combo.get("name",""); cmap = combo.get("map",{})
                    if not cmap:
                        prog.log(f"({idx}) BỎ QUA '{name}': map rỗng."); prog.step(1); continue
                    prog.set_status(f"({idx}/{len(self.combos)}) '{name}' (#{number})")

                    cmb = robot_create_combo_with_number(cases, number, name_hint=name)
                    if cmb is None:
                        prog.log(f"   -> LỖI CreateCombination #{number}"); prog.step(1); continue
                    try: cmb.Name = name
                    except Exception: pass
                    try: fac = cmb.CaseFactors
                    except Exception: fac = None

                    added = 0
                    if fac is not None:
                        try: fac.BeginMultiOperation()
                        except Exception: pass
                        for cid_str, coef in cmap.items():
                            try:
                                cnum = int(cid_str); cf = float(coef)
                            except Exception:
                                continue
                            ok=False
                            try: fac.New(cnum, cf); ok=True
                            except Exception:
                                try:
                                    f = fac.New(); f.CaseNumber = cnum; f.Factor = cf; ok=True
                                except Exception: pass
                            if ok: added += 1
                        try: fac.EndMultiOperation()
                        except Exception: pass
                    prog.log(f"   -> Thêm {added} factors.")
                    created += 1
                    prog.step(1, hint=f"Đã tạo '{name}'")
            finally:
                try: cases.EndMultiOperation()
                except Exception: pass
                try: rob.Project.ViewMngr.Refresh()
                except Exception: pass

            prog.set_status("Hoàn tất."); prog.log(f">> Tạo thành công: {created}/{len(self.combos)}")
            prog.close()
            try: rob.Visible = True
            except Exception: pass
            messagebox.showinfo("Thành công", f"Đã tạo {created} tổ hợp (ROBOT).")
            return

        # -------- SAP2000 --------
        sap_model, _ = connect_sap2000(visible=not self.fast_mode.get())
        if not sap_model:
            prog.log("!! Không tìm thấy phiên SAP2000."); prog.close()
            messagebox.showerror("Không kết nối", "Không tìm thấy phiên SAP2000 đang mở.")
            return

        if self.fast_mode.get() and not ENGINE_ALREADY_CLEAN:
            names_set = {c["name"] for c in self.combos if c.get("name")}
            if names_set:
                prog.log(f"- Đang xóa trước combo trùng tên ({len(names_set)}) ...")
            removed, tried_surface = sap_delete_combos_by_names(sap_model, names_set, prog=prog if names_set else None)
            surf = "LoadCombinations/RespCombo" if tried_surface else "unknown"
            prog.log(f"- Đã xóa trước {removed} combo trùng tên (surface={surf}).")

        resolver = self._get_case_name_resolver(ctx)
        created_total = 0
        for idx, combo in enumerate(self.combos, start=1):
            if prog.cancelled: break
            name = combo.get("name",""); cmap = combo.get("map",{})
            prog.set_status(f"({idx}/{len(self.combos)}) '{name}'")
            added = sap_create_combo_and_add_factors(sap_model, name, cmap, resolver)
            created_total += 1 if added>0 else 0
            prog.log(f"   -> Thêm {added} factors.")
            prog.step(1, hint=f"Đã tạo '{name}'")

        try:
            ui = getattr(sap_model, "View", None)
            if ui and hasattr(ui, "RefreshView"):
                ui.RefreshView(0, False)
        except Exception: pass

        prog.set_status("Hoàn tất."); prog.log(f">> Tạo xong {created_total}/{len(self.combos)} tổ hợp (SAP2000)."); prog.close()
        messagebox.showinfo("Thành công", f"Đã tạo {created_total} tổ hợp (SAP2000).")

    def delete_all_engine_combos(self):
        from robot_helper import iterate_cases, pick
        target = self.push_engine.get().upper()
        prog = ProgressDialog(self, title=f"Xóa combo trong {target}", total=100)
        prog.set_status("Đang dò danh sách...")

        if target == "ROBOT":
            rob, _ = connect_robot()
            if not rob:
                prog.close(); messagebox.showerror("Không kết nối", "Không tìm thấy phiên Robot đang mở."); return
            cases = rob.Project.Structure.Cases
            to_del = []
            try:
                for c in iterate_cases(cases):
                    try:
                        _ = c.CaseFactors
                        n = pick(c, "Number","No","Id","ID","Index")
                        if isinstance(n, int): to_del.append(n)
                    except Exception: continue
            except Exception: pass

            if not to_del:
                prog.close(); messagebox.showinfo("Xóa combo", "Không tìm thấy combo nào (ROBOT)."); return

            if not messagebox.askyesno("Xóa hết combo", f"Tìm thấy {len(to_del)} combo (ROBOT).\nXóa TẤT CẢ không?", icon="warning"):
                prog.close(); return

            prog.set_total(len(to_del)); prog.log(f">> Sẽ xóa {len(to_del)} combo (ROBOT).")
            try: cases.BeginMultiOperation()
            except Exception: pass

            deleted = 0
            for i, n in enumerate(sorted(to_del, reverse=True), start=1):
                if prog.cancelled: break
                try:
                    cases.Delete(n); deleted += 1
                except Exception:
                    try:
                        cases.DeleteCase(n); deleted += 1
                    except Exception: pass
                prog.step(1, hint=f"Xóa combo #{n} ({i}/{len(to_del)})")

            try: cases.EndMultiOperation()
            except Exception: pass
            try: rob.Project.ViewMngr.Refresh()
            except Exception: pass

            prog.close()
            messagebox.showinfo("Xóa combo", f"Đã xóa {deleted}/{len(to_del)} combo (ROBOT).")
            return

        # ----- SAP2000 -----
        from sap_helper import sap_delete_combos_by_names
        sap_model, _ = connect_sap2000()
        if not sap_model:
            prog.close(); messagebox.showerror("Không kết nối", "Không tìm thấy phiên SAP2000 đang mở."); return

        names_all = set()
        try:
            lc = sap_model.LoadCombinations
            ret, n, c_names = lc.GetNameList()
            for nm in (c_names or []): names_all.add(nm)
        except Exception: pass
        try:
            rc = sap_model.RespCombo
            ret, n, c_names = rc.GetNameList()
            for nm in (c_names or []): names_all.add(nm)
        except Exception: pass

        if not names_all:
            prog.close(); messagebox.showinfo("Xóa combo", "Không tìm thấy combo nào (SAP2000)."); return

        if not messagebox.askyesno("Xóa hết combo", f"Tìm thấy {len(names_all)} combo (SAP2000).\nXóa TẤT CẢ không?", icon="warning"):
            prog.close(); return

        prog.set_total(len(names_all))
        prog.log(f">> Sẽ xóa {len(names_all)} combo (SAP2000).")

        deleted = 0
        try:
            lc = sap_model.LoadCombinations
            for i, nm in enumerate(sorted(names_all), start=1):
                if prog.cancelled: break
                try:
                    ok = lc.Delete(nm)
                    if ok in (0, None): deleted += 1
                except Exception:
                    try:
                        lc.SetSelected(nm, True); lc.DeleteSelected(); deleted += 1
                    except Exception: pass
                prog.step(1, hint=f"Xóa '{nm}' ({i}/{len(names_all)})")
        except Exception: pass

        try:
            rc = sap_model.RespCombo
            for nm in list(names_all):
                try:
                    r = rc.Delete(nm)
                    if r in (0, None): deleted += 1
                except Exception: pass
        except Exception: pass

        prog.close()
        messagebox.showinfo("Xóa combo", f"Đã xóa {deleted}/{len(names_all)} combo (SAP2000).")
        return